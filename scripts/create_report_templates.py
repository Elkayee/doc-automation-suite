"""Build the editable DOCX style templates; never touches user workspaces."""

import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Cm

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.core.config import TemplateConfig
from src.core.docx_helpers import DocxHelpers
from src.core.report_fields import configure_styles


def create_ptit_cover(source, destination, *, individual=False):
    """Adapt the supplied cover without rebuilding its table or logo."""
    document = Document(source)
    existing = {style.style_id for style in document.styles}
    for style in Document().styles:
        if style.style_id not in existing:
            document.styles.element.append(deepcopy(style.element))
    table = document.tables[0]
    for row in table.rows:
        for paragraph in row.cells[0].paragraphs:
            if paragraph.paragraph_format.first_line_indent is None:
                paragraph.paragraph_format.first_line_indent = Cm(0)

    def replace_text(paragraph, text):
        nodes = paragraph._p.xpath('.//w:t')
        if nodes:
            nodes[0].text = text
            for node in nodes[1:]:
                node.text = ''
        else:
            paragraph.add_run(text)

    title = table.cell(2, 0).paragraphs
    replace_text(title[1], 'BÀI TẬP CÁ NHÂN' if individual else 'BÀI TẬP NHÓM')
    replace_text(title[2], 'HỌC PHẦN {{subject}}')
    replace_text(title[3], '{{title}}')
    details = table.cell(3, 0).paragraphs
    replace_text(details[0], 'Giảng viên hướng dẫn: {{supervisor}}')
    replace_text(details[1], 'Lớp học phần: {{class}}')
    if individual:
        replace_text(details[2], 'Sinh viên thực hiện: {{author}}')
        replace_text(details[3], 'Mã sinh viên: {{student_id}}')
        for paragraph in details[4:9]:
            replace_text(paragraph, '')
    else:
        replace_text(details[2], 'Nhóm thực hiện: {{group}}')
        for index, paragraph in enumerate(details[4:9], 1):
            replace_text(paragraph, '{{member_' + str(index) + '}}')
    replace_text(table.cell(4, 0).paragraphs[3], '{{date}}')
    document.save(destination)


def main():
    templates = Path(__file__).resolve().parents[1] / 'templates'
    for name in ['bao_cao_bai_tap_lon', 'bao_cao_cong_so', 'bien_ban']:
        project = templates / name
        config = TemplateConfig.load(project / 'config.yaml')
        document = Document()
        configure_styles(document, config)
        DocxHelpers.apply_page_settings(document, DocxHelpers.get_page_settings(config))
        document.save(project / config.docx_template)
        print(name)
    source = templates / 'tieu_luan_nd30/Cover_Ptit.docx'
    for name, individual in [('tieu_luan_nd30', False), ('tieu_luan_ca_nhan', True)]:
        create_ptit_cover(source, templates / name / 'template.docx', individual=individual)
        print(name)


if __name__ == '__main__':
    main()
