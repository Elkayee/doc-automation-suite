import re

from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from src.core.markdown_semantics import text_parts


def add_field(paragraph, instruction, cached=''):
    for kind in ('begin', 'instruction', 'separate', 'result', 'end'):
        run = OxmlElement('w:r')
        if kind in ('begin', 'separate', 'end'):
            element = OxmlElement('w:fldChar')
            element.set(qn('w:fldCharType'), kind)
            if kind == 'begin':
                element.set(qn('w:dirty'), 'true')
        else:
            element = OxmlElement('w:instrText' if kind == 'instruction' else 'w:t')
            element.set(qn('xml:space'), 'preserve')
            element.text = instruction if kind == 'instruction' else cached
        run.append(element)
        paragraph._p.append(run)


def start_bookmark(paragraph, identifier, number):
    if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_]{0,39}', identifier):
        raise ValueError(f'Mã tham chiếu không hợp lệ: {identifier}')
    element = OxmlElement('w:bookmarkStart')
    element.set(qn('w:id'), str(number))
    element.set(qn('w:name'), identifier)
    paragraph._p.append(element)


def end_bookmark(paragraph, number):
    element = OxmlElement('w:bookmarkEnd')
    element.set(qn('w:id'), str(number))
    paragraph._p.append(element)


def add_caption(doc, label, text, index, identifier=''):
    paragraph = doc.add_paragraph(style='Caption')
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.keep_with_next = label == 'Bảng'
    if identifier:
        start_bookmark(paragraph, identifier, index + (10000 if label == 'Bảng' else 20000))
    paragraph.add_run(label + ' ')
    add_field(paragraph, f'SEQ {label} \\* ARABIC', str(index))
    if identifier:
        end_bookmark(paragraph, index + (10000 if label == 'Bảng' else 20000))
    paragraph.add_run(': ' + text)
    return paragraph


def configure_styles(doc, config):
    settings = config.settings if config else {}
    font_name = settings.get('font_name', 'Times New Roman')
    font_size = float(settings.get('font_size', 14))
    for name in ['Normal', 'Caption', 'Title', 'Subtitle'] + [f'Heading {i}' for i in range(1, 10)]:
        style = doc.styles[name]
        style.font.name = font_name
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font_name)
        if name.startswith('Heading'):
            level = int(name.split()[-1])
            style.font.size = Pt(max(font_size, font_size + 4 - level * 2))
            style.font.bold = True
            style.paragraph_format.first_line_indent = Cm(0)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True
        elif name == 'Caption':
            style.font.size = Pt(max(10, font_size - 2))
            style.paragraph_format.first_line_indent = Cm(0)
    update = doc.settings.element.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields')
        doc.settings.element.append(update)
    update.set(qn('w:val'), 'true')
    if config:
        doc.core_properties.title = config.metadata.get('title') or config.name
        doc.core_properties.author = config.metadata.get('author') or ''


def set_page_numbering(section, start=1, fmt='decimal'):
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    add_field(paragraph, 'PAGE', str(start))
    page = section._sectPr.find(qn('w:pgNumType'))
    if page is None:
        page = OxmlElement('w:pgNumType')
        section._sectPr.append(page)
    page.set(qn('w:start'), str(start))
    page.set(qn('w:fmt'), fmt)


def set_report_header(section, config):
    if not config:
        return
    text = config.settings.get('header_text', config.metadata.get('organization') or '')
    section.header.is_linked_to_previous = False
    paragraph = section.header.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if text:
        paragraph.add_run(str(text)).font.size = Pt(10)


def start_body_section(doc, config=None):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_page_numbering(section)
    set_report_header(section, config)
    return section


def add_report_cover(doc, config):
    metadata = config.metadata
    for key in ('organization', 'department'):
        value = metadata.get(key)
        if value:
            p = doc.add_paragraph(str(value))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
    doc.add_paragraph()
    title = doc.add_paragraph(metadata.get('title') or '[Chưa nhập tên báo cáo]', style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.space_before = Pt(36)
    title.paragraph_format.space_after = Pt(36)
    for key, label in config.metadata_fields.items():
        if key in {'title', 'organization', 'department'}:
            continue
        value = metadata.get(key)
        if value:
            p = doc.add_paragraph(f'{label}: {value}')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_page_numbering(section, fmt='lowerRoman')


def add_reference_runs(paragraph, text, references, write_text):
    if '[[REF:' not in text:
        write_text(paragraph, text)
        return
    for code, part in text_parts(text):
        if code:
            write_text(paragraph, part)
        else:
            _add_reference_text(paragraph, part, references, write_text)


def _add_reference_text(paragraph, text, references, write_text):
    position = 0
    for match in re.finditer(r'\[\[REF:\s*([A-Za-z][A-Za-z0-9_]*)\s*\]\]', text):
        write_text(paragraph, text[position:match.start()])
        identifier = match.group(1)
        if identifier not in references:
            raise ValueError(f'Tham chiếu chưa tồn tại: {identifier}')
        add_field(paragraph, f'REF {identifier} \\h', references[identifier])
        position = match.end()
    write_text(paragraph, text[position:])
