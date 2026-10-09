import csv
import json
from pathlib import Path

import yaml
from openpyxl import load_workbook

from src.core.config import TemplateConfig
from src.core.export import compile_document
from src.core.template_manager import TemplateManager

CONTROL_COLUMNS = {'workspace', 'template', 'docx_out', 'md_out', 'pdf_out', 'cache_dir', 'formats', 'mode', 'engine'}


def _row_to_job(row):
    job = {key: value for key, value in row.items() if key in CONTROL_COLUMNS and value not in (None, '')}
    job['metadata'] = {key: value if value not in (None, '') else None for key, value in row.items() if key not in CONTROL_COLUMNS}
    if 'formats' in job:
        job['formats'] = [part.strip() for part in str(job['formats']).replace(';', ',').split(',') if part.strip()]
    return job


def load_batch_manifest(path: Path):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as handle:
            reader = csv.DictReader(handle)
            headers = reader.fieldnames or []
            if not headers or len(set(headers)) != len(headers) or any(not h for h in headers):
                raise ValueError('CSV cần tên cột rõ ràng, không trùng')
            rows = list(reader)
            if any(None in row for row in rows):
                raise ValueError('CSV có dòng nhiều giá trị hơn số cột')
        jobs = [_row_to_job(row) for row in rows if any(value not in (None, '') for value in row.values())]
    elif path.suffix.lower() == '.xlsx':
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = iter(workbook.active.iter_rows(values_only=True))
            headers = list(next(rows, ()))
            if not headers or len(set(headers)) != len(headers) or any(not isinstance(h, str) or not h for h in headers):
                raise ValueError('XLSX cần tên cột rõ ràng, không trùng ở dòng đầu')
            jobs = [_row_to_job(dict(zip(headers, row, strict=True))) for row in rows if any(value is not None for value in row)]
        finally:
            workbook.close()
    else:
        text = path.read_text(encoding='utf-8')
        data = json.loads(text) if path.suffix.lower() == '.json' else yaml.safe_load(text)
        jobs = data.get('jobs') if isinstance(data, dict) else data
    if not isinstance(jobs, list) or not jobs or any(not isinstance(job, dict) for job in jobs):
        raise ValueError('Manifest cần danh sách jobs không rỗng')
    return jobs


def run_batch(jobs, *, base_dir: Path, templates_dir: Path | None = None):
    base_dir = Path(base_dir).resolve()
    results = []
    used_outputs = set()
    protected_roots = []
    protected_files = set()
    for item in jobs:
        name = item.get('workspace')
        if not isinstance(name, str):
            continue
        project = (base_dir / name).resolve()
        if not project.is_relative_to(base_dir):
            continue
        protected_roots.extend(project / directory for directory in ['chapters', 'assets', '.recovery'])
        protected_files.add(project / 'config.yaml')
        protected_files.add(project / 'template.docx')
        if (project / 'config.yaml').is_file():
            try:
                config = TemplateConfig.load(project / 'config.yaml')
                protected_files.update((project / path).resolve() for path in [config.docx_template, config.bibliography, config.csl] if path)
            except (ValueError, OSError):
                pass
    for index, item in enumerate(jobs, 1):
        workspace = item.get('workspace')
        try:
            if '_error' in item:
                raise ValueError(item['_error'])
            if not isinstance(workspace, str) or not workspace.strip():
                raise ValueError('Thiếu workspace')
            project = (base_dir / workspace).resolve()
            if not project.is_relative_to(base_dir):
                raise ValueError('Workspace của batch phải nằm trong thư mục gốc của manifest')
            options = {key: value for key, value in item.items() if key not in {'workspace', 'template'}}
            for key in ('docx_out', 'md_out', 'pdf_out', 'cache_dir'):
                if options.get(key):
                    options[key] = (base_dir / options[key]).resolve()
                    if not options[key].is_relative_to(base_dir):
                        raise ValueError(f'{key} nằm ngoài phạm vi batch')
            destination = options.get('docx_out') or project / 'build' / f'{project.name}.docx'
            destination = Path(destination).resolve()
            if options.get('docx_out') and not options.get('md_out'):
                options['md_out'] = destination.with_suffix('.md')
            item_outputs = {destination, destination.with_suffix('.export.json'),
                            Path(options.get('md_out') or project / 'build' / 'assembled.md').resolve()}
            if 'pdf' in options.get('formats', ['docx']):
                item_outputs.add(Path(options.get('pdf_out') or destination.with_suffix('.pdf')).resolve())
            if item_outputs & used_outputs:
                raise ValueError('Đầu ra bị trùng trong batch; mỗi dòng cần tên file riêng')
            checked_paths = item_outputs | ({options['cache_dir']} if options.get('cache_dir') else set())
            if any(path in protected_files or any(path.is_relative_to(root) for root in protected_roots) for path in checked_paths):
                raise ValueError('Đầu ra/cache không được ghi đè đầu vào của các mục trong batch')
            used_outputs.update(item_outputs)
            if item.get('template'):
                if templates_dir is None:
                    raise ValueError('Chưa cấu hình thư mục templates')
                TemplateManager(templates_dir).create_project(item['template'], project)
            result = compile_document(project, **options)
            results.append({'index': index, **result.to_dict()})
        except Exception as exc:
            results.append({'index': index, 'workspace': workspace, 'success': False, 'error': str(exc)})
    failed = sum(not result['success'] for result in results)
    return {'success': failed == 0, 'total': len(results), 'succeeded': len(results) - failed, 'failed': failed, 'results': results}
