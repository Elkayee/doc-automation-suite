import hashlib
import json
import re
import shutil
import tempfile
import unicodedata
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from pypdf import PdfReader

from src.core.assembler import DocumentAssembler
from src.core.docx_builder import DocxBuilder
from src.core.docx_helpers import DocxHelpers
from src.core.file_io import atomic_write, export_lock, file_hash
from src.core.markdown_image import build_markdown_image, parse_markdown_image_line
from src.core.markdown_semantics import CITATION_RE, content_lines, has_academic_markup, text_parts
from src.core.renderers import refresh_and_render_pdf


class ExportValidationError(ValueError):
    def __init__(self, errors):
        self.errors = list(errors)
        super().__init__('\n'.join(self.errors))


@dataclass
class ExportResult:
    success: bool
    mode: str
    engine: str
    workspace: str
    assembled_markdown: str
    compiled_docx: str
    compiled_pdf: str | None
    chapters_processed: int
    warnings: list[str]
    field_status: str
    input_sha256: str
    output_sha256: dict[str, str]

    def to_dict(self):
        return asdict(self)


def project_issues(config, entries, workspace):
    issues = []
    if config.docx_template and not (workspace / config.docx_template).is_file():
        issues.append(f'Thiếu DOCX mẫu: {config.docx_template}')
    for key in config.required_metadata:
        if not str(config.metadata.get(key) or '').strip():
            issues.append(f'Chưa nhập {config.metadata_fields.get(key, key)}')
    contents = {entry.filename: entry.content for entry in entries}
    for filename in config.required_files:
        if not contents.get(filename, '').strip():
            issues.append(f'Chương bắt buộc trống: {filename}')
    for entry in entries:
        for number, line in enumerate(entry.content.splitlines(), 1):
            for key, value in config.metadata.items():
                if value is not None:
                    line = line.replace('{{' + key + '}}', value)
            if '[Nhập nội dung vào đây]' in line or re.search(r'\{\{[^}]+\}\}', line):
                issues.append(f'{entry.filename}:{number}: còn nội dung chưa điền')
    return issues


def _copy_input(source, snapshot, workspace, hashes):
    source = Path(source)
    if not source.resolve().is_relative_to(workspace):
        raise ExportValidationError([f'Đầu vào nằm ngoài dự án: {source.name}'])
    relative = source.relative_to(workspace)
    target = snapshot / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)
    hashes[relative.as_posix()] = file_hash(target)


def _validate_docx(path):
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None or 'word/document.xml' not in archive.namelist():
            raise ExportValidationError(['DOCX không hợp lệ'])
    Document(path)


def _normalized_text(text):
    return ''.join(unicodedata.normalize('NFKC', text).replace('\u00ad', '').split())


def _validate_pdf(path, title=None, expected_docx=None):
    reader = PdfReader(path, strict=True)
    if not reader.pages:
        raise ExportValidationError(['PDF không có trang'])
    text = '\n'.join(page.extract_text() or '' for page in reader.pages)
    normalized = re.sub(r'\s+', ' ', text).strip()
    if not normalized:
        raise ExportValidationError(['PDF không có văn bản có thể đọc/tìm'])
    if title and re.sub(r'\s+', ' ', title).strip() not in normalized:
        raise ExportValidationError(['PDF thiếu tên báo cáo'])
    if 'Cap nhat Muc luc trong Word bang Update Field.' in text:
        raise ExportValidationError(['Mục lục chưa được cập nhật trong PDF'])
    if expected_docx:
        document = Document(expected_docx)
        expected = []
        for paragraph in document.element.xpath('.//w:p'):
            instructions = paragraph.xpath('.//w:instrText')
            if any((instruction.text or '').strip().startswith('TOC ') for instruction in instructions):
                continue
            expected.extend(node.text for node in paragraph.xpath('.//w:t') if node.text and node.text.strip())
        normalized_pdf = _normalized_text(text)
        missing = [value for value in expected if _normalized_text(value) not in normalized_pdf]
        if missing:
            raise ExportValidationError(['PDF thiếu nội dung: ' + value[:120] for value in missing[:10]])
        if document.inline_shapes and not any(page.images for page in reader.pages):
            raise ExportValidationError(['PDF thiếu hình ảnh của tài liệu'])


def publish_outputs(outputs):
    previous = {target: target.read_bytes() if target.is_file() else None for target in outputs}
    published = []
    try:
        for target, data in outputs.items():
            atomic_write(target, data)
            published.append(target)
    except Exception:
        for target in reversed(published):
            if previous[target] is None:
                target.unlink(missing_ok=True)
            else:
                atomic_write(target, previous[target])
        raise


def compile_document(
    workspace_dir: Path, *, docx_out: Path | None = None, md_out: Path | None = None,
    pdf_out: Path | None = None, formats=('docx',), mode='draft', engine='auto',
    cache_dir: Path | None = None, metadata: dict | None = None,
) -> ExportResult:
    workspace = Path(workspace_dir).resolve()
    if not workspace.is_dir() or not (workspace / 'build').resolve().is_relative_to(workspace):
        raise ValueError('Dự án không tồn tại hoặc build nằm ngoài dự án')
    with export_lock(workspace / 'build' / '.export.lock'):
        return _compile_document(workspace, docx_out=docx_out, md_out=md_out, pdf_out=pdf_out,
                                 formats=formats, mode=mode, engine=engine, cache_dir=cache_dir, metadata=metadata)


def _compile_document(
    workspace_dir: Path, *, docx_out: Path | None = None, md_out: Path | None = None,
    pdf_out: Path | None = None, formats=('docx',), mode='draft', engine='auto',
    cache_dir: Path | None = None, metadata: dict | None = None,
) -> ExportResult:
    workspace = Path(workspace_dir).resolve()
    if mode not in {'draft', 'final'} or engine not in {'auto', 'builtin', 'academic'}:
        raise ValueError('Chế độ xuất không hợp lệ')
    if not formats or any(item not in {'docx', 'pdf'} for item in formats):
        raise ValueError('Định dạng xuất phải là docx hoặc pdf')
    assembler = DocumentAssembler(workspace)
    _content, entries = assembler.assemble_with_metadata()
    config = assembler.get_config()
    if metadata is not None:
        config.metadata.update({str(key): str(value) if value is not None else None for key, value in metadata.items()})
    build = workspace / 'build'
    if not build.resolve().is_relative_to(workspace):
        raise ValueError('Thư mục build nằm ngoài dự án')
    build.mkdir(exist_ok=True)
    final_docx = Path(docx_out or build / f'{workspace.name}.docx').resolve()
    final_md = Path(md_out or build / 'assembled.md').resolve()
    final_pdf = Path(pdf_out or final_docx.with_suffix('.pdf')).resolve() if 'pdf' in formats else None
    manifest_path = final_docx.with_suffix('.export.json')
    destinations = [final_docx, final_md, manifest_path] + ([final_pdf] if final_pdf else [])
    lock_path = (build / '.export.lock').resolve()
    if lock_path in destinations:
        raise ValueError('Đầu ra không được thay thế khóa xuất của dự án')
    if len(set(destinations)) != len(destinations):
        raise ValueError('Các file đầu ra phải có đường dẫn khác nhau')
    sources = {entry.path.resolve() for entry in entries} | {assembler.config_path.resolve()}
    sources.update((workspace / p).resolve() for p in [config.docx_template, config.bibliography, config.csl] if p)
    if any(target in sources or target.is_relative_to(workspace / 'chapters') or target.is_relative_to(workspace / 'assets')
           or target.is_relative_to(workspace / '.recovery') for target in destinations):
        raise ValueError('Đầu ra không được ghi đè cấu hình hoặc chương nguồn')

    with tempfile.TemporaryDirectory(prefix='.export-', dir=build) as temporary:
        stage = Path(temporary).resolve()
        assert stage.is_relative_to(build.resolve())
        snapshot = stage / 'project'
        snapshot.mkdir()
        hashes = {}
        for entry in entries:
            _copy_input(entry.path, snapshot, workspace, hashes)
        _copy_input(workspace / 'config.yaml', snapshot, workspace, hashes)
        extras = [config.docx_template, config.bibliography, config.csl]
        for relative in filter(None, extras):
            source = workspace / relative
            if source.is_file():
                _copy_input(source, snapshot, workspace, hashes)
        assets = workspace / 'assets'
        if assets.exists():
            for source in assets.rglob('*'):
                if source.is_file():
                    _copy_input(source, snapshot, workspace, hashes)
        for entry in entries:
            for line in content_lines(entry.content):
                image = parse_markdown_image_line(line)
                if image:
                    source = DocxHelpers.resolve_media_path(workspace, entry.path, image.path)
                    if source.is_file():
                        _copy_input(source, snapshot, workspace, hashes)
        for relative, digest in hashes.items():
            if file_hash(workspace / relative) != digest:
                raise ExportValidationError(['Đầu vào đã thay đổi khi bắt đầu xuất. Hãy xuất lại sau khi lưu.'])
        config.save(snapshot / 'config.yaml')
        staged_assembler = DocumentAssembler(snapshot)
        _content, frozen_entries = staged_assembler.assemble_with_metadata()
        warnings = project_issues(config, frozen_entries, snapshot)
        if mode == 'final' and warnings:
            raise ExportValidationError(warnings)
        lines = []
        for entry in frozen_entries:
            lines.append(f'<!-- FILE: {entry.filename} -->')
            for line in entry.content.splitlines():
                image = parse_markdown_image_line(line)
                if image:
                    source = DocxHelpers.resolve_media_path(workspace, workspace / 'chapters' / entry.filename, image.path)
                    if source.is_file():
                        line = build_markdown_image(source.relative_to(workspace).as_posix(), alt=image.alt,
                                                   caption=image.caption, width=image.width, align=image.align,
                                                   identifier=image.identifier)
                for key, value in config.metadata.items():
                    if value is not None:
                        line = line.replace('{{' + key + '}}', value)
                lines.append(line)
            lines.append('')
        markdown = '\n'.join(lines)
        selected_engine = engine
        if selected_engine == 'auto':
            selected_engine = 'academic' if config.bibliography or has_academic_markup(markdown) else 'builtin'
        academic = None
        cache = Path(cache_dir or build / 'render_cache')
        resolved_cache = cache.resolve()
        if resolved_cache == lock_path:
            raise ValueError('Cache không được thay thế khóa xuất của dự án')
        if (resolved_cache == workspace or resolved_cache.is_relative_to(workspace / 'chapters')
                or resolved_cache.is_relative_to(workspace / 'assets')
                or resolved_cache.is_relative_to(workspace / '.recovery')
                or any(source.is_relative_to(resolved_cache) for source in sources)):
            raise ValueError('Cache không được ghi vào vùng chứa đầu vào hoặc bản phục hồi')
        cache.mkdir(parents=True, exist_ok=True)
        if selected_engine == 'academic':
            from src.core.academic import AcademicRenderer

            academic = AcademicRenderer(snapshot, config, cache)
            markdown, citation_warnings = academic.render_citations(markdown)
            warnings.extend(citation_warnings)
        elif any(CITATION_RE.search(part) for line in content_lines(markdown) for code, part in text_parts(line) if not code):
            warnings.append('Trích dẫn chưa được xử lý; chọn chế độ học thuật hoặc tự động')
        staged_md = snapshot / 'assembled.md'
        atomic_write(staged_md, markdown)
        builder = DocxBuilder(snapshot)
        builder.academic = academic
        builder.build_from_markdown(str(staged_md), cache)
        warnings.extend(builder.warnings)
        warnings = list(dict.fromkeys(warnings))
        if mode == 'final' and warnings:
            raise ExportValidationError(warnings)
        staged_docx = stage / 'composed.docx'
        builder.save(staged_docx)
        _validate_docx(staged_docx)
        with zipfile.ZipFile(staged_docx) as package:
            has_fields = False
            for name in package.namelist():
                if name.startswith('word/') and name.endswith('.xml'):
                    tree = etree.fromstring(package.read(name), parser=etree.XMLParser(resolve_entities=False, no_network=True))
                    if tree.find('.//' + qn('w:instrText')) is not None or tree.find('.//' + qn('w:fldSimple')) is not None:
                        has_fields = True
                        break
        field_status = 'needs_update' if has_fields else 'not_required'
        staged_pdf = None
        if 'pdf' in formats or (mode == 'final' and has_fields):
            expected_docx = staged_docx
            staged_docx, staged_pdf = refresh_and_render_pdf(staged_docx, stage / 'render')
            _validate_docx(staged_docx)
            _validate_pdf(staged_pdf, config.metadata.get('title'), expected_docx)
            refreshed_document = Document(staged_docx)
            if 'Cap nhat Muc luc trong Word bang Update Field.' in ''.join(refreshed_document.element.itertext()):
                raise ExportValidationError(['DOCX còn mục lục chưa cập nhật'])
            field_status = 'updated'
        outputs = {final_docx: staged_docx.read_bytes(), final_md: markdown.encode('utf-8')}
        if final_pdf:
            outputs[final_pdf] = staged_pdf.read_bytes()
        input_data = {'files': hashes, 'config': config.model_dump()}
        input_digest = hashlib.sha256(json.dumps(input_data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        result = ExportResult(
            True, mode, selected_engine, str(workspace), str(final_md), str(final_docx),
            str(final_pdf) if final_pdf else None, len(frozen_entries), warnings, field_status,
            input_digest, {str(target): hashlib.sha256(data).hexdigest() for target, data in outputs.items()},
        )
        outputs[manifest_path] = json.dumps(result.to_dict(), ensure_ascii=False, indent=2).encode()
        publish_outputs(outputs)
        return result
