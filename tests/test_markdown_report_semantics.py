from zipfile import ZipFile

import pytest
from docx import Document

from src.core.config import TemplateConfig
from src.core.export import ExportValidationError, compile_document
from src.core.markdown_semantics import has_academic_markup


def project(tmp_path, text):
    (tmp_path / 'chapters').mkdir()
    TemplateConfig(required_files=['Ch01.md'], settings={'numbered_captions': True}).save(tmp_path / 'config.yaml')
    Document().save(tmp_path / 'template.docx')
    (tmp_path / 'chapters/Ch01.md').write_text(text, encoding='utf-8')
    return tmp_path


def test_code_examples_and_currency_do_not_select_academic_engine():
    assert not has_academic_markup('Giá $500 và $700. Email [me@example.com].\n`[@code]`\n```md\n$code$\n[@code]\n```')
    assert has_academic_markup('Công thức $E=mc^2$ và [xem @source].')


def test_fenced_caption_examples_do_not_create_phantom_references(tmp_path):
    root = project(tmp_path, '# Nội dung\n\n```md\n[[TABLE: results | Ví dụ cú pháp]]\n```\n\n'
                   '[[TABLE: results | Bảng thật]]\n| Cột | Giá trị |\n| --- | --- |\n| A | 1 |\n\n'
                   'Xem [[REF: results]]. Mã `[[REF: missing]]` giữ nguyên.\n')
    result = compile_document(root)
    paragraphs = [p.text for p in Document(result.compiled_docx).paragraphs]
    assert any('Xem Bảng 1.' in p and '[[REF: missing]]' in p for p in paragraphs)
    with ZipFile(result.compiled_docx) as package:
        assert package.read('word/document.xml').count(b'REF results') == 1


def test_final_builtin_cannot_report_unresolved_citations_as_success(tmp_path):
    root = project(tmp_path, '# Nội dung\n\nNguồn: [@source].\n')
    with pytest.raises(ExportValidationError, match='Trích dẫn chưa được xử lý'):
        compile_document(root, mode='final', engine='builtin')


def test_escaped_pipe_and_code_pipe_stay_in_the_same_table_cell(tmp_path):
    root = project(tmp_path, '# Bảng\n\n| Nội dung | Giá trị |\n| --- | --- |\n| A\\|B | `x|y` |\n')
    result = compile_document(root)
    table = Document(result.compiled_docx).tables[0]
    assert len(table.columns) == 2
    assert [cell.text for cell in table.rows[1].cells] == ['A|B', 'x|y']
