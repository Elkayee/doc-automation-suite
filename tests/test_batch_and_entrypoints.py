import json
from zipfile import ZipFile

import pytest
from click.testing import CliRunner
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

import src.api as api
from src.cli import cli
from src.core.batch import load_batch_manifest, run_batch
from src.core.config import TemplateConfig
from src.core.export import compile_document
from src.core.file_io import export_lock


def make_project(path, *, require_amount=False):
    (path / 'chapters').mkdir(parents=True)
    config = TemplateConfig(required_files=['Ch01.md'], required_metadata=['amount'] if require_amount else [])
    config.save(path / 'config.yaml')
    Document().save(path / 'template.docx')
    (path / 'chapters' / 'Ch01.md').write_text('# Báo cáo\n\nGiá trị: {{amount}}\n' if require_amount else '# Báo cáo\n\nĐã đủ nội dung.\n',
                                            encoding='utf-8')
    return path


def test_mixed_batch_keeps_success_after_failure_and_rejects_duplicate_output(tmp_path):
    first = make_project(tmp_path / 'first')
    second = make_project(tmp_path / 'second')
    result = run_batch([{'workspace': 'first'}, {'workspace': 'missing'}, {'workspace': 'second'},
                        {'workspace': 'first'}], base_dir=tmp_path)
    assert result['succeeded'] == 2 and result['failed'] == 2
    assert [item['success'] for item in result['results']] == [True, False, True, False]
    assert (first / 'build/first.docx').is_file() and (second / 'build/second.docx').is_file()


@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
def test_metadata_rows_preserve_missing_and_zero_and_final_checks_apply(tmp_path, extension):
    make_project(tmp_path / 'project', require_amount=True)
    source = tmp_path / f'rows.{extension}'
    if extension == 'csv':
        source.write_text('workspace,docx_out,mode,amount\nproject,zero.docx,final,0\nproject,missing.docx,final,\n', encoding='utf-8')
    else:
        workbook = Workbook()
        workbook.active.append(['workspace', 'docx_out', 'mode', 'amount'])
        workbook.active.append(['project', 'zero.docx', 'final', 0])
        workbook.active.append(['project', 'missing.docx', 'final', None])
        workbook.save(source)
    jobs = load_batch_manifest(source)
    assert jobs[1]['metadata']['amount'] is None
    assert jobs[0]['metadata']['amount'] in ('0', 0)
    result = run_batch(jobs, base_dir=tmp_path)
    assert result['results'][0]['success'] and not result['results'][1]['success']
    assert not (tmp_path / 'missing.docx').exists()
    assert 'Giá trị: 0' in '\n'.join(p.text for p in Document(tmp_path / 'zero.docx').paragraphs)
    assert TemplateConfig.load(tmp_path / 'project/config.yaml').metadata == {}


def test_cli_api_core_share_semantic_output_and_machine_result(tmp_path, monkeypatch):
    project = make_project(tmp_path / 'workspaces' / 'example')
    monkeypatch.setattr(api, 'BASE_DIR', tmp_path)
    core = compile_document(project, docx_out=project / 'core.docx')
    invoked = CliRunner().invoke(cli, ['compile', str(project), '--docx-out', str(project / 'cli.docx'), '--json'])
    assert invoked.exit_code == 0, invoked.output
    cli_result = json.loads(invoked.stdout)
    response = TestClient(api.app).post('/workspaces/compile', json={'workspace_name': 'example', 'docx_out': 'api.docx'})
    assert response.status_code == 200
    paths = [core.compiled_docx, cli_result['compiled_docx'], response.json()['compiled_docx']]
    xml = []
    for path in paths:
        with ZipFile(path) as archive:
            xml.append(archive.read('word/document.xml'))
    assert xml[0] == xml[1] == xml[2]


def test_api_batch_isolated_invalid_path_and_final_validation(tmp_path, monkeypatch):
    make_project(tmp_path / 'workspaces' / 'example', require_amount=True)
    monkeypatch.setattr(api, 'BASE_DIR', tmp_path)
    client = TestClient(api.app)
    response = client.post('/workspaces/batch', json={'jobs': [
        {'workspace_name': '../outside'},
        {'workspace_name': 'example', 'mode': 'final', 'metadata': {'amount': 0}},
    ]})
    assert response.status_code == 200
    assert response.json()['succeeded'] == 1 and response.json()['failed'] == 1
    assert not (tmp_path / 'outside').exists()
    failed = client.post('/workspaces/compile', json={'workspace_name': 'example', 'mode': 'final'})
    assert failed.status_code == 422
    outside = client.post('/workspaces/compile', json={'workspace_name': 'example', 'docx_out': '../outside.docx'})
    assert outside.status_code == 400


def test_export_lock_blocks_second_process_handle_and_releases_after_failure(tmp_path):
    path = tmp_path / 'lock'
    with export_lock(path):
        with pytest.raises(RuntimeError, match='tác vụ khác'):
            with export_lock(path):
                pass
    with export_lock(path):
        pass


def test_batch_rejects_scope_escape(tmp_path):
    make_project(tmp_path / 'one')
    make_project(tmp_path / 'two')
    result = run_batch([{'workspace': '../outside', 'template': 'bao_cao_cong_so'},
                        {'workspace': 'one', 'docx_out': '../outside.docx'}], base_dir=tmp_path)
    assert result['failed'] == 2
    assert not (tmp_path.parent / 'outside').exists()
    assert not (tmp_path.parent / 'outside.docx').exists()


def test_batch_rejects_a_shared_pdf_output_before_second_export(tmp_path, monkeypatch):
    from types import SimpleNamespace

    calls = []

    def export(workspace, **options):
        calls.append(workspace)
        return SimpleNamespace(to_dict=lambda: {'success': True})

    monkeypatch.setattr('src.core.batch.compile_document', export)
    result = run_batch([{'workspace': 'one', 'formats': ['pdf'], 'pdf_out': 'shared.pdf'},
                        {'workspace': 'two', 'formats': ['pdf'], 'pdf_out': 'shared.pdf'}], base_dir=tmp_path)
    assert result['succeeded'] == 1 and result['failed'] == 1
    assert len(calls) == 1


def test_two_metadata_rows_have_independent_markdown_artifacts(tmp_path):
    make_project(tmp_path / 'project', require_amount=True)
    result = run_batch([{'workspace': 'project', 'docx_out': 'one.docx', 'metadata': {'amount': 1}},
                        {'workspace': 'project', 'docx_out': 'two.docx', 'metadata': {'amount': 2}}], base_dir=tmp_path)
    assert result['succeeded'] == 2
    assert '{{amount}}' not in (tmp_path / 'one.md').read_text(encoding='utf-8')
    assert 'Giá trị: 1' in (tmp_path / 'one.md').read_text(encoding='utf-8')
    assert 'Giá trị: 2' in (tmp_path / 'two.md').read_text(encoding='utf-8')


def test_batch_cannot_overwrite_another_jobs_chapter(tmp_path):
    make_project(tmp_path / 'one')
    make_project(tmp_path / 'two')
    source = tmp_path / 'two/chapters/Ch01.md'
    original = source.read_bytes()
    result = run_batch([{'workspace': 'one', 'docx_out': 'two/chapters/Ch01.md'},
                        {'workspace': 'two'}], base_dir=tmp_path)
    assert result['failed'] == 1 and result['succeeded'] == 1
    assert source.read_bytes() == original
