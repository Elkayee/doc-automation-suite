from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from src.core.export import ExportValidationError, _validate_pdf


def text_pdf(path, text):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(('BT /F1 12 Tf 40 750 Td (' + text + ') Tj ET').encode('ascii'))
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(path)


def test_title_only_pdf_cannot_pass_for_a_document_with_missing_body(tmp_path):
    expected = Document()
    expected.add_paragraph('Report')
    expected.add_paragraph('Required chapter content and caption')
    source = tmp_path / 'expected.docx'
    expected.save(source)
    pdf = tmp_path / 'wrong.pdf'
    text_pdf(pdf, 'Report')
    with pytest.raises(ExportValidationError, match='PDF thiếu nội dung'):
        _validate_pdf(pdf, 'Report', source)


def test_pdf_with_correct_text_but_no_required_image_is_rejected(tmp_path):
    expected = Document()
    expected.add_paragraph('Report')
    expected.add_picture(str(Path(__file__).resolve().parents[1] / 'test_extracted.png'))
    source = tmp_path / 'expected.docx'
    expected.save(source)
    pdf = tmp_path / 'wrong.pdf'
    text_pdf(pdf, 'Report')
    with pytest.raises(ExportValidationError, match='PDF thiếu hình ảnh'):
        _validate_pdf(pdf, 'Report', source)
