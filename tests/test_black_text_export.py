import hashlib
from zipfile import ZipFile

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from src.core.config import TemplateConfig
from src.core.docx_builder import DocxBuilder


def test_export_blackens_template_text_and_theme_colors_without_editing_source(tmp_path):
    document = Document()
    document.styles['Title'].font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    colored = document.add_paragraph().add_run('Nội dung có định dạng')
    colored.bold = True
    colored.font.size = Pt(16)
    colored.font.color.rgb = RGBColor(255, 0, 0)
    color = colored._r.rPr.color
    color.set(qn('w:themeColor'), 'accent1')
    color.set(qn('w:themeTint'), '80')
    for paragraph in [
        document.add_heading('Tiêu đề', 2),
        document.sections[0].header.paragraphs[0],
        document.sections[0].footer.paragraphs[0],
        document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0],
    ]:
        paragraph.add_run('Chữ trong mẫu').font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), document.part.relate_to('https://example.com', RT.HYPERLINK, is_external=True))
    link_run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    style = OxmlElement('w:rStyle')
    style.set(qn('w:val'), 'Hyperlink')
    properties.append(style)
    link_run.append(properties)
    text = OxmlElement('w:t')
    text.text = 'Liên kết'
    link_run.append(text)
    hyperlink.append(link_run)
    document.paragraphs[0]._p.append(hyperlink)
    footnotes = b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:rPr><w:color w:val="00FF00"/></w:rPr><w:t>Footnote text</w:t></w:r></w:p></w:footnote></w:footnotes>'
    raw_part = Part(PackURI('/word/footnotes.xml'), CT.WML_FOOTNOTES, footnotes, document.part.package)
    document.part.relate_to(raw_part, RT.FOOTNOTES)
    template = tmp_path / 'template.docx'
    document.save(template)
    before = hashlib.sha256(template.read_bytes()).hexdigest()
    TemplateConfig(docx_template='template.docx').save(tmp_path / 'config.yaml')
    output = tmp_path / 'report.docx'
    DocxBuilder(tmp_path).save(output)
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before
    reopened = Document(output)
    assert reopened.paragraphs[0].runs[0].bold
    assert reopened.paragraphs[0].runs[0].font.size == Pt(16)
    assert 'Nội dung có định dạng' in reopened.paragraphs[0].text
    assert reopened.part.rels[hyperlink.get(qn('r:id'))].target_ref == 'https://example.com'
    with ZipFile(output) as package:
        for name in package.namelist():
            if not name.startswith('word/') or not name.endswith('.xml'):
                continue
            root = parse_xml(package.read(name))
            for run in root.iter(qn('w:r')):
                color = run.find(qn('w:rPr')).find(qn('w:color'))
                assert color is not None, (name, run.xml)
                assert dict(color.attrib) == {qn('w:val'): '000000'}, name
            for color in root.iter(qn('w:color')):
                if color.getparent().tag == qn('w:rPr'):
                    assert dict(color.attrib) == {qn('w:val'): '000000'}, name
        assert b'Footnote text' in package.read('word/footnotes.xml')
