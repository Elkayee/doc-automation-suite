from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn

from src.core.config import TemplateConfig
from src.core.docx_builder import DocxBuilder
from src.core.export import ExportValidationError, compile_document
from src.core.report_fields import fill_template_cover
from src.core.template_manager import TemplateManager

TEMPLATES = Path(__file__).resolve().parents[1] / 'templates'
SOURCE = TEMPLATES / 'tieu_luan_nd30/Cover_Ptit.docx'


@pytest.mark.parametrize(
    ('template', 'label', 'authors', 'absent'),
    [
        (
            'tieu_luan_nd30',
            'BÀI TẬP NHÓM',
            {'group': '07', 'member_1': 'An - B22DCCN001', 'member_2': 'Bình - B22DCCN002'},
            'Sinh viên thực hiện:',
        ),
        (
            'tieu_luan_ca_nhan',
            'BÀI TẬP CÁ NHÂN',
            {'author': 'Lan', 'student_id': 'B22DCCN003'},
            'Danh sách thành viên:',
        ),
    ],
)
def test_ptit_cover_contains_correct_authors_and_preserves_source_layout(tmp_path, template, label, authors, absent):
    TemplateManager(TEMPLATES).create_project(template, tmp_path)
    config = TemplateConfig.load(tmp_path / 'config.yaml')
    config.metadata.update(
        title='Đề tài kiểm chứng',
        supervisor='Giảng viên kiểm chứng',
        **{'class': 'SKD1103-01', 'date': 'Hà Nội, tháng 10/2026'},
        **authors,
    )
    config.save(tmp_path / 'config.yaml')
    before = (tmp_path / 'template.docx').read_bytes()
    source_before = SOURCE.read_bytes()
    result = compile_document(tmp_path)
    document = Document(result.compiled_docx)
    assert 'Bìa PTIT lấy từ' not in ''.join(document.element.itertext())
    assert len(document.tables) == 1
    assert sum(node.text == 'Đề tài kiểm chứng' for node in document.element.xpath('.//w:t')) == 1
    cover = document.tables[0]
    text = '\n'.join(row.cells[0].text for row in cover.rows)
    assert text.count(label) == 1
    assert absent not in text
    assert '{{' not in text and 'Chưa nhập' not in text
    assert '(số nhóm)' not in text and '(Họ và tên)' not in text
    for value in authors.values():
        assert value in text
    assert 'Đề tài kiểm chứng' in text
    assert len(cover.rows) == 5
    original = Document(SOURCE)
    assert len(cover._tbl.xpath('.//w:drawing | .//w:pict')) == len(
        original.tables[0]._tbl.xpath('.//w:drawing | .//w:pict')
    )
    for row, source_row in zip(cover.rows, original.tables[0].rows):
        for paragraph, source_paragraph in zip(row.cells[0].paragraphs, source_row.cells[0].paragraphs):
            assert paragraph.paragraph_format.first_line_indent == (
                source_paragraph.paragraph_format.first_line_indent or 0
            )
    assert cover._tbl.tblPr.xml == original.tables[0]._tbl.tblPr.xml
    assert (
        document.sections[0]._sectPr.find(qn('w:pgMar')).attrib
        == original.sections[0]._sectPr.find(qn('w:pgMar')).attrib
    )
    assert document.sections[1].start_type == original.sections[0].start_type
    with ZipFile(SOURCE) as source, ZipFile(result.compiled_docx) as output:
        assert output.read('word/media/image1.emf') == source.read('word/media/image1.emf')
    assert (tmp_path / 'template.docx').read_bytes() == before
    assert SOURCE.read_bytes() == source_before


def test_cover_placeholders_split_across_runs_preserve_formatting_and_literal_values():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Name: {{au').bold = True
    paragraph.add_run('thor}} / {{student_id}} suffix').italic = True
    config = TemplateConfig(metadata={'author': 'Lan {{literal}}', 'student_id': '001'})
    fill_template_cover(document, config)
    assert paragraph.text == 'Name: Lan {{literal}} / 001 suffix'
    assert paragraph.runs[0].bold and paragraph.runs[1].italic


def test_missing_individual_metadata_warns_and_final_failure_preserves_previous_export(tmp_path):
    TemplateManager(TEMPLATES).create_project('tieu_luan_ca_nhan', tmp_path)
    draft = compile_document(tmp_path)
    assert any('Mã sinh viên' in warning for warning in draft.warnings)
    document = Document(draft.compiled_docx)
    assert '[Chưa nhập Mã sinh viên]' in document.tables[0].cell(3, 0).text
    before = Path(draft.compiled_docx).read_bytes()
    with pytest.raises(ExportValidationError, match='Mã sinh viên'):
        compile_document(tmp_path, mode='final')
    assert Path(draft.compiled_docx).read_bytes() == before


def test_existing_non_ptit_template_keeps_its_original_behavior(tmp_path):
    document = Document()
    document.add_paragraph('Unchanged {{author}}')
    document.save(tmp_path / 'template.docx')
    TemplateConfig(metadata={'author': 'Lan'}).save(tmp_path / 'config.yaml')
    builder = DocxBuilder(tmp_path)
    assert builder.doc.paragraphs[0].text == 'Unchanged {{author}}'
    assert len(builder.doc.sections) == 1
