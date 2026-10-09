import json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.shared import Cm

from src.core.config import TemplateConfig
from src.core.export import ExportValidationError, compile_document, publish_outputs
from src.core.file_io import ExternalEditError, atomic_write, file_hash, save_chapter


@pytest.fixture
def project(tmp_path):
    (tmp_path / 'chapters').mkdir()
    config = TemplateConfig(required_files=['Ch01.md'])
    config.save(tmp_path / 'config.yaml')
    Document().save(tmp_path / 'template.docx')
    (tmp_path / 'chapters' / 'Ch01.md').write_text('# Báo cáo\n\nNội dung đã hoàn thành.\n', encoding='utf-8')
    return tmp_path


def test_missing_image_is_warning_in_draft_and_rejected_in_final_without_replacing_output(project):
    chapter = project / 'chapters' / 'Ch01.md'
    original = compile_document(project, mode='final')
    previous = Path(original.compiled_docx).read_bytes()
    chapter.write_text('# Báo cáo\n\n![Minh chứng](missing.png)\n', encoding='utf-8')
    with pytest.raises(ExportValidationError, match='Thiếu ảnh'):
        compile_document(project, mode='final')
    assert Path(original.compiled_docx).read_bytes() == previous
    draft = compile_document(project)
    assert draft.success and any('Thiếu ảnh' in warning for warning in draft.warnings)


def test_output_cannot_replace_template_or_manifest_and_markdown_cannot_collide(project):
    with pytest.raises(ValueError):
        compile_document(project, docx_out=project / 'template.docx')
    with pytest.raises(ValueError):
        compile_document(project, docx_out=project / 'out.docx', md_out=project / 'out.export.json')


@pytest.mark.parametrize('option', ['docx_out', 'md_out', 'cache_dir'])
def test_output_cannot_replace_the_live_export_lock(project, option):
    with pytest.raises(ValueError, match='khóa xuất'):
        compile_document(project, **{option: project / 'build/.export.lock'})


def test_existing_template_keeps_later_landscape_sections(project):
    template = Document()
    second = template.add_section()
    second.orientation = WD_ORIENTATION.LANDSCAPE
    second.page_width, second.page_height = Cm(29.7), Cm(21)
    template.save(project / 'template.docx')
    result = compile_document(project)
    document = Document(result.compiled_docx)
    assert document.sections[1].page_width > document.sections[1].page_height


def test_atomic_write_failure_preserves_original_and_removes_temporary_file(tmp_path):
    target = tmp_path / 'chapter.md'
    target.write_text('original', encoding='utf-8')
    with patch('src.core.file_io.os.replace', side_effect=PermissionError('locked')):
        with pytest.raises(PermissionError):
            atomic_write(target, 'replacement')
    assert target.read_text(encoding='utf-8') == 'original'
    assert list(tmp_path.glob('*.tmp')) == []


def test_external_edit_preserves_external_file_and_local_recovery_draft(project):
    chapter = project / 'chapters' / 'Ch01.md'
    expected = file_hash(chapter)
    chapter.write_text('external edit', encoding='utf-8')
    with pytest.raises(ExternalEditError):
        save_chapter(chapter, 'my unsaved changes', expected)
    assert chapter.read_text(encoding='utf-8') == 'external edit'
    assert (project / '.recovery' / 'drafts' / chapter.name).read_text(encoding='utf-8') == 'my unsaved changes'


def test_failure_publishing_second_artifact_rolls_back_first(tmp_path):
    first, second = tmp_path / 'report.docx', tmp_path / 'report.pdf'
    first.write_bytes(b'old docx')
    second.write_bytes(b'old pdf')
    actual = atomic_write

    def injected(path, content):
        if path == second:
            raise PermissionError('locked pdf')
        return actual(path, content)

    with patch('src.core.export.atomic_write', side_effect=injected):
        with pytest.raises(PermissionError):
            publish_outputs({first: b'new docx', second: b'new pdf'})
    assert first.read_bytes() == b'old docx' and second.read_bytes() == b'old pdf'


def test_semantic_fields_captions_refs_and_page_sections(project):
    config = TemplateConfig.load(project / 'config.yaml')
    config.settings.update({'report_fields': True, 'numbered_captions': True})
    config.metadata = {'title': 'Báo cáo mẫu', 'author': 'Người kiểm thử'}
    config.metadata_fields = {'title': 'Đề tài', 'author': 'Người lập'}
    config.save(project / 'config.yaml')
    (project / 'chapters' / 'Ch01.md').write_text(
        '[[COVER]]\n\n[[TOC]]\n\n[[BODY]]\n\n# 1. Nội dung {#intro}\n\n'
        'Xem [[REF: table1]].\n\n[[TABLE: table1 | Kết quả thử nghiệm]]\n'
        '| Cột | Giá trị |\n| --- | --- |\n| A | 1 |\n\n[[LANDSCAPE]]\n\n## 1.1. Phụ lục\n',
        encoding='utf-8',
    )
    result = compile_document(project)
    assert result.field_status == 'needs_update'
    with ZipFile(result.compiled_docx) as archive:
        xml = archive.read('word/document.xml').decode()
        assert 'SEQ Bảng' in xml and 'REF table1' in xml and 'TOC ' in xml
        assert 'bookmarkStart' in xml and 'w:name="table1"' in xml
        assert 'lowerRoman' in xml and 'landscape' in xml
    document = Document(result.compiled_docx)
    assert len(document.sections) == 4
    assert any(p.style.name == 'Heading 1' and 'Nội dung' in p.text for p in document.paragraphs)
    manifest = json.loads(Path(result.compiled_docx).with_suffix('.export.json').read_text(encoding='utf-8'))
    assert manifest['input_sha256'] and manifest['output_sha256'][result.compiled_docx] == file_hash(Path(result.compiled_docx))
